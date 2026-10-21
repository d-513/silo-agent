"""Themed 16:9 decks with python-pptx, plus a proof you can look at.

    import sys; sys.path.insert(0, "/opt/silo/skills/presentations/scripts")
    from slides import Deck
    d = Deck("ink")                       # ink | garden | ember (see DESIGN.md)
    d.title("Q3 operations review", "Northwind Logistics · October 2026")
    d.section("01", "Delivery")
    d.points("On-time delivery is back above target", ["New carrier took 38% of regional parcels", "..."])
    d.stats("The quarter in numbers", [("96.2%", "On-time delivery"), ("€1.84m", "Warehouse cost")])
    d.chart("Delivery recovered every month", "bot/chart.png", note="Share of parcels on time")
    d.image("The new Lyon site", "bot/site.jpg", caption="Opened 2 September")
    d.table("Cost by site", [["Site", "Q2", "Q3"], ["Lyon", "€812k", "€706k"]], numeric=(1, 2))
    d.quote("The best quarter since 2019.", "Head of operations")
    d.end("Questions", "ops@northwind.example")
    d.save("/workspace/deck.pptx")        # also renders bot/proof/deck-sheet.png and prints warnings

Each method takes notes="…" for speaker notes. Charts: make a PNG with theme.mpl() (native
pptx charts show as placeholders in the Silo console), 16:9-ish, about 11 x 4.6 in.
"""
import html
import os
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import theme as _theme  # noqa: E402

W, H = 13.333, 7.5
M = 0.8  # side margin, inches


class Deck:
    def __init__(self, theme="ink"):
        self.t = _theme.get(theme)
        self.slides = []  # each: dict(bg=..., els=[...], notes=...)
        self.warnings = []

    # ---- building blocks -------------------------------------------------
    def _slide(self, bg=None, notes=None):
        s = {"bg": bg or self.t.paper, "els": [], "notes": notes}
        self.slides.append(s)
        return s

    def _text(self, s, x, y, w, h, paras, anchor="top", align="left"):
        s["els"].append({"k": "text", "box": (x, y, w, h), "paras": paras, "anchor": anchor, "align": align})

    def _rect(self, s, x, y, w, h, fill):
        s["els"].append({"k": "rect", "box": (x, y, w, h), "fill": fill})

    def _p(self, text, size, color=None, bold=False, italic=False, font=None, after=0, caps=False):
        return {"text": text.upper() if caps else text, "size": size, "color": color or self.t.ink, "bold": bold,
                "italic": italic, "font": font or self.t.office_body, "after": after, "caps": caps}

    def _head(self, s, title, kicker=None):
        n = len(self.slides)
        if len(title) > 70:
            self.warnings.append(f"slide {n}: title is {len(title)} chars; say it in under 70")
        if kicker:
            self._text(s, M, 0.55, W - 2 * M, 0.35, [self._p(kicker, 11, self.t.accent, bold=True, caps=True)])
        self._text(s, M, 0.9, W - 2 * M, 1.1, [self._p(title, 34, font=self.t.office_display)])

    # ---- slide kinds -----------------------------------------------------
    def title(self, title, subtitle="", notes=None):
        s = self._slide(self.t.accent, notes)
        self._text(s, M, 2.2, W - 2 * M - 1.5, 2.6, [self._p(title, 46, self.t.paper, font=self.t.office_display)], anchor="bottom")
        self._rect(s, M, 5.05, 1.2, 0.04, self.t.paper)
        if subtitle:
            self._text(s, M, 5.3, W - 2 * M, 0.8, [self._p(subtitle, 16, self.t.accent_soft)])

    def section(self, number, title, notes=None):
        s = self._slide(self.t.tint, notes)
        self._text(s, M, 2.1, 3, 1.4, [self._p(number, 66, self.t.accent, font=self.t.office_display)], anchor="bottom")
        self._text(s, M, 3.6, W - 2 * M, 1.4, [self._p(title, 40, font=self.t.office_display)])

    def points(self, title, items, kicker=None, notes=None):
        n = len(self.slides) + 1
        if len(items) > 5:
            self.warnings.append(f"slide {n}: {len(items)} bullets; keep 3-5, split the rest onto another slide")
        for it in items:
            if len(it) > 110:
                self.warnings.append(f"slide {n}: a bullet is {len(it)} chars; one line of thought each")
        s = self._slide(None, notes)
        self._head(s, title, kicker)
        paras = []
        for it in items:
            paras.append(self._p("—  " + it, 22, after=16))
        self._rect(s, M, 2.25, 0.6, 0.04, self.t.accent)
        self._text(s, M, 2.55, W - 2 * M - 2.0, H - 3.2, paras)

    def stats(self, title, figures, kicker=None, notes=None):
        n = len(self.slides) + 1
        if len(figures) > 4:
            self.warnings.append(f"slide {n}: {len(figures)} figures; 2-4 read at a glance")
        s = self._slide(None, notes)
        self._head(s, title, kicker)
        k = max(1, len(figures))
        gap = 0.45
        cw = (W - 2 * M - gap * (k - 1)) / k
        for i, f in enumerate(figures):
            value, label = f[0], f[1]
            x = M + i * (cw + gap)
            self._rect(s, x, 2.7, cw, 0.05, self.t.accent)
            paras = [self._p(value, 60, font=self.t.office_display, after=4), self._p(label, 14, self.t.muted)]
            if len(f) > 2:
                paras.append(self._p(f[2], 13, self.t.accent, after=0))
            self._text(s, x, 2.9, cw, 2.6, paras)

    def chart(self, title, png, note="", kicker=None, notes=None):
        s = self._slide(None, notes)
        self._head(s, title, kicker)
        self._image(s, png, M, 2.15, W - 2 * M, 4.55 if note else 4.9)
        if note:
            self._text(s, M, 6.8, W - 2 * M, 0.4, [self._p(note, 11, self.t.muted)])

    def image(self, title, path, caption="", kicker=None, notes=None):
        s = self._slide(None, notes)
        self._image(s, path, M + 5.0, 0.8, W - 2 * M - 5.0, H - 1.6)
        self._text(s, M, 0.8, 4.4, 0.4, [self._p(kicker or "", 11, self.t.accent, bold=True, caps=True)])
        self._text(s, M, 1.25, 4.4, 3.5, [self._p(title, 34, font=self.t.office_display)])
        if caption:
            self._text(s, M, 5.2, 4.4, 1.5, [self._p(caption, 13, self.t.muted)], anchor="bottom")

    def table(self, title, rows, numeric=(), kicker=None, notes=None):
        n = len(self.slides) + 1
        if len(rows) > 8:
            self.warnings.append(f"slide {n}: table has {len(rows) - 1} rows; keep 7 or fewer, or put it in a handout")
        s = self._slide(None, notes)
        self._head(s, title, kicker)
        s["els"].append({"k": "table", "box": (M, 2.3, W - 2 * M, 0.5 * len(rows)), "rows": rows, "numeric": tuple(numeric)})

    def quote(self, text, who="", notes=None):
        s = self._slide(self.t.tint, notes)
        self._text(s, M + 0.6, 1.4, W - 2 * M - 1.8, 3.8, [self._p("“" + text + "”", 36, font=self.t.office_display, italic=True)], anchor="bottom")
        if who:
            self._text(s, M + 0.6, 5.5, W - 2 * M, 0.5, [self._p("— " + who, 15, self.t.muted)])

    def end(self, title="Thank you", line="", notes=None):
        s = self._slide(self.t.ink, notes)
        self._text(s, M, 2.4, W - 2 * M, 1.6, [self._p(title, 44, self.t.paper, font=self.t.office_display)], anchor="bottom")
        if line:
            self._text(s, M, 4.2, W - 2 * M, 0.6, [self._p(line, 16, self.t.accent_soft)])

    def _image(self, s, path, x, y, w, h):
        from PIL import Image
        iw, ih = Image.open(path).size
        r = min(w / iw, h / ih)
        dw, dh = iw * r, ih * r
        s["els"].append({"k": "image", "box": (x + (w - dw) / 2, y + (h - dh) / 2, dw, dh), "path": str(path)})

    # ---- output ----------------------------------------------------------
    def save(self, out, proof=True):
        self._pptx(out)
        if not proof:
            return
        stem = Path(out).stem
        proof_dir = Path("/workspace/bot/proof") if Path("/workspace").is_dir() else Path("bot/proof")
        proof_dir.mkdir(parents=True, exist_ok=True)
        page = proof_dir / f"{stem}.html"
        page.write_text(self._html())
        print(f"pptx: {out} ({len(self.slides)} slides)", flush=True)
        render = Path(__file__).resolve().parent / "render.py"
        subprocess.run([sys.executable, str(render), str(page), str(proof_dir / f"{stem}-proof.pdf"), "--slides"])
        print("(the proof approximates PowerPoint: same layout and sizes, the box's own fonts)")
        for w in self.warnings:
            print("WARNING: " + w)

    def _pptx(self, out):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt

        rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(W), Inches(H)
        blank = prs.slide_layouts[6]
        for sd in self.slides:
            sl = prs.slides.add_slide(blank)
            bg = sl.background.fill
            bg.solid()
            bg.fore_color.rgb = rgb(sd["bg"])
            for e in sd["els"]:
                x, y, w, h = (Inches(v) for v in e["box"])
                if e["k"] == "rect":
                    sh = sl.shapes.add_shape(1, x, y, w, h)
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = rgb(e["fill"])
                    sh.line.fill.background()
                    sh.shadow.inherit = False
                elif e["k"] == "image":
                    sl.shapes.add_picture(e["path"], x, y, w, h)
                elif e["k"] == "text":
                    tf = sl.shapes.add_textbox(x, y, w, h).text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = MSO_ANCHOR.BOTTOM if e["anchor"] == "bottom" else MSO_ANCHOR.TOP
                    for i, p in enumerate(e["paras"]):
                        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        para.alignment = PP_ALIGN.LEFT
                        para.space_after = Pt(p["after"])
                        para.line_spacing = 1.05 if p["size"] >= 28 else 1.15
                        r = para.add_run()
                        r.text = p["text"]
                        f = r.font
                        f.size, f.bold, f.italic, f.name = Pt(p["size"]), p["bold"], p["italic"], p["font"]
                        f.color.rgb = rgb(p["color"])
                elif e["k"] == "table":
                    rows = e["rows"]
                    gt = sl.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
                    tblpr = gt._tbl.tblPr
                    tblpr.set("firstRow", "0")
                    tblpr.set("bandRow", "0")
                    for ri, row in enumerate(rows):
                        for ci, val in enumerate(row):
                            c = gt.cell(ri, ci)
                            c.fill.solid()
                            c.fill.fore_color.rgb = rgb(self.t.tint if ri == 0 else self.t.paper)
                            c.text = str(val)
                            para = c.text_frame.paragraphs[0]
                            para.alignment = PP_ALIGN.RIGHT if ci in e["numeric"] else PP_ALIGN.LEFT
                            for r in para.runs:
                                r.font.size = Pt(14 if ri else 12)
                                r.font.bold = ri == 0
                                r.font.name = self.t.office_body
                                r.font.color.rgb = rgb(self.t.muted if ri == 0 else self.t.ink)
            if sd["notes"]:
                sl.notes_slide.notes_text_frame.text = sd["notes"]
        prs.save(out)

    def _html(self):
        t = self.t
        fam = lambda f: f"'{f}', " + ("'Linux Libertine O', serif" if f == "Georgia" else "'Inter', sans-serif")
        out = [f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: {W}in {H}in; margin: 0; }}
body {{ margin: 0; }}
.s {{ position: relative; width: {W}in; height: {H}in; overflow: hidden; break-after: page; }}
.e {{ position: absolute; }}
.t {{ display: flex; flex-direction: column; }}
.t p {{ margin: 0; }}
table {{ border-collapse: collapse; width: 100%; font-family: {fam(t.office_body)}; }}
td {{ padding: 0 .12in; height: .5in; border-bottom: 1pt solid {t.rule}; }}
</style></head><body>"""]
        for sd in self.slides:
            out.append(f'<section class="s" style="background-color:{sd["bg"]}">')
            for e in sd["els"]:
                x, y, w, h = e["box"]
                pos = f"left:{x}in;top:{y}in;width:{w}in;"
                if e["k"] == "rect":
                    out.append(f'<div class="e" style="{pos}height:{h}in;background-color:{e["fill"]}"></div>')
                elif e["k"] == "image":
                    out.append(f'<img class="e" style="{pos}height:{h}in" src="{Path(e["path"]).resolve().as_uri()}">')
                elif e["k"] == "text":
                    just = "flex-end" if e["anchor"] == "bottom" else "flex-start"
                    ps = []
                    for p in e["paras"]:
                        lh = 1.05 if p["size"] >= 28 else 1.15
                        ps.append(f'<p style="font-family:{fam(p["font"])};font-size:{p["size"]}pt;line-height:{lh};color:{p["color"]};'
                                  f'font-weight:{600 if p["bold"] else 400};font-style:{"italic" if p["italic"] else "normal"};'
                                  f'letter-spacing:{".08em" if p["caps"] else "0"};margin-bottom:{p["after"]}pt">{html.escape(p["text"])}</p>')
                    out.append(f'<div class="e t" style="{pos}height:{h}in;justify-content:{just}">{"".join(ps)}</div>')
                elif e["k"] == "table":
                    trs = []
                    for ri, row in enumerate(e["rows"]):
                        tds = "".join(
                            f'<td style="text-align:{"right" if ci in e["numeric"] else "left"};font-size:{14 if ri else 12}pt;'
                            f'font-weight:{600 if ri == 0 else 400};color:{t.muted if ri == 0 else t.ink};'
                            f'background-color:{t.tint if ri == 0 else t.paper}">{html.escape(str(v))}</td>'
                            for ci, v in enumerate(row))
                        trs.append(f"<tr>{tds}</tr>")
                    out.append(f'<table class="e" style="{pos}">{"".join(trs)}</table>')
            out.append("</section>")
        out.append("</body></html>")
        return "\n".join(out)
